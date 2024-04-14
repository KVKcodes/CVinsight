from django.shortcuts import render
from django.http import HttpResponse
from .forms import CVUploadForm
import os
import docx
import re
from phonenumbers import parse, is_valid_number
from openpyxl import Workbook
import PyPDF2
import aspose.words as aw

def extract_text_from_pdf(pdf_file):
    text = ''
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    for page_num in range(len(pdf_reader.pages) ):
        page = pdf_reader.pages[page_num] 
        text += page.extract_text()
    return text

def extract_info_from_cv(cv_file):
    text = ''
    email = ''
    phone = ''
    
    # Extract text from docx
    if cv_file.name.endswith('.docx'):
        doc = docx.Document(cv_file)
        for paragraph in doc.paragraphs:
            text += paragraph.text + '\n'

    # Extract text from pdf
    elif cv_file.name.endswith('.pdf'):
        text = extract_text_from_pdf(cv_file)
        
    else:
        # Handle unsupported file types here
        pass

    
    # Extract email and phone
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'[\+\(]?[1-9][0-9 .\-\(\)]{8,}[0-9]'
    matches_email = re.findall(email_pattern, text)
    matches_phone = re.findall(phone_pattern, text)
    if matches_email:
        email = matches_email[0]
    if matches_phone:
        phone = matches_phone[0]
    
    return text, email, phone

def generate_xls(data):
    wb = Workbook()
    ws = wb.active
    ws.append(['Text', 'Email', 'Phone'])
    for item in data:
        ws.append(item)
    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = 'attachment; filename="cv_data.xls"'
    wb.save(response)
    return response

def upload_cv(request):
    if request.method == 'POST':
        form = CVUploadForm(request.POST, request.FILES)
        if form.is_valid():
            cv_file = form.cleaned_data['file']
            text, email, phone = extract_info_from_cv(cv_file)
            data = [(text, email, phone)]
            return generate_xls(data)
    else:
        form = CVUploadForm()
    return render(request, 'upload_cv.html', {'form': form})

